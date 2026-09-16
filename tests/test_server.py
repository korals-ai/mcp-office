"""Tests for the MCP server wiring.

Verifies the server constructs and registers the ``convert`` tool, and that
the tool maps conversion failures to an MCP error rather than crashing. The
office service behind ``convert`` is the ``convert_stub`` fixture; the real
one is exercised by the post-deploy integration suite.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from src import server


@pytest.mark.asyncio
async def test_convert_tool_is_registered() -> None:
    tools = await server.mcp.list_tools()
    names = {t.name for t in tools}
    assert "convert" in names


@pytest.mark.asyncio
async def test_pdf_to_images_tool_is_registered() -> None:
    tools = await server.mcp.list_tools()
    names = {t.name for t in tools}
    assert "pdf_to_images" in names


@pytest.mark.asyncio
async def test_pdf_to_images_tool_describes_src_and_dpi() -> None:
    tools = await server.mcp.list_tools()
    tool = next(t for t in tools if t.name == "pdf_to_images")
    schema = tool.inputSchema
    assert "src" in schema["properties"]
    assert "dpi" in schema["properties"]
    assert "src" in schema.get("required", [])


@pytest.mark.asyncio
async def test_pdf_extract_text_tool_is_registered() -> None:
    tools = await server.mcp.list_tools()
    names = {t.name for t in tools}
    assert "pdf_extract_text" in names


@pytest.mark.parametrize("name", ["author_xlsx", "author_docx", "author_pptx", "author_pdf"])
@pytest.mark.asyncio
async def test_author_tools_are_registered(name: str) -> None:
    tools = await server.mcp.list_tools()
    assert name in {t.name for t in tools}


@pytest.mark.asyncio
async def test_author_xlsx_tool_requires_path_and_rows() -> None:
    tools = await server.mcp.list_tools()
    schema = next(t for t in tools if t.name == "author_xlsx").inputSchema
    assert "path" in schema["properties"]
    assert "rows" in schema["properties"]
    required = set(schema.get("required", []))
    assert {"path", "rows"} <= required


@pytest.mark.asyncio
async def test_pdf_extract_text_tool_describes_src() -> None:
    tools = await server.mcp.list_tools()
    tool = next(t for t in tools if t.name == "pdf_extract_text")
    schema = tool.inputSchema
    assert "src" in schema["properties"]
    assert "src" in schema.get("required", [])


@pytest.mark.asyncio
async def test_convert_tool_describes_src_and_to() -> None:
    tools = await server.mcp.list_tools()
    convert_tool = next(t for t in tools if t.name == "convert")
    schema = convert_tool.inputSchema
    assert "src" in schema["properties"]
    assert "to" in schema["properties"]
    # src is required, to has a default and is optional.
    assert "src" in schema.get("required", [])


def test_convert_missing_source_raises(tmp_path) -> None:
    # The tool wrapper delegates to office_convert and lets the error
    # propagate (FastMCP renders it as an MCP tool error).
    from src.office_convert import OfficeConvertError

    with pytest.raises(OfficeConvertError):
        server.convert(str(tmp_path / "absent.docx"), "pdf")


def test_convert_url_comes_from_the_environment_with_no_default() -> None:
    # OFFICE_CONVERT_URL is read once at import (the runner exports it); the
    # tool wrapper passes exactly that base to office_convert.
    import os

    assert os.environ["OFFICE_CONVERT_URL"] == server.CONVERT_URL


def test_convert_tool_posts_to_the_configured_service(tmp_path, convert_stub, monkeypatch) -> None:
    monkeypatch.setattr(server, "CONVERT_URL", convert_stub.url)
    src = tmp_path / "a.docx"
    src.write_bytes(b"PK\x03\x04 x")
    convert_stub.body = b"%PDF-1.7 ok"
    out = server.convert(str(src), "pdf")
    assert out == str(tmp_path / "a.pdf")
    assert convert_stub.requests[0]["path"] == "/cool/convert-to/pdf"


def test_pdf_to_images_missing_source_raises(tmp_path) -> None:
    from src.pdf_rasterize import PdfRasterizeError

    with pytest.raises(PdfRasterizeError):
        server.pdf_to_images(str(tmp_path / "absent.pdf"))


def test_pdf_extract_text_missing_source_raises(tmp_path) -> None:
    from src.pdf_text import PdfTextExtractError

    with pytest.raises(PdfTextExtractError):
        server.pdf_extract_text(str(tmp_path / "absent.pdf"))


@pytest.mark.asyncio
async def test_xlsx_extract_cells_tool_is_registered() -> None:
    tools = await server.mcp.list_tools()
    names = {t.name for t in tools}
    assert "xlsx_extract_cells" in names


@pytest.mark.asyncio
async def test_xlsx_extract_cells_tool_describes_src_and_sheet() -> None:
    tools = await server.mcp.list_tools()
    tool = next(t for t in tools if t.name == "xlsx_extract_cells")
    schema = tool.inputSchema
    assert "src" in schema["properties"]
    assert "sheet" in schema["properties"]
    assert "src" in schema.get("required", [])


def test_xlsx_extract_cells_missing_source_raises(tmp_path) -> None:
    from src.xlsx_extract import XlsxExtractError

    with pytest.raises(XlsxExtractError):
        server.xlsx_extract_cells(str(tmp_path / "absent.xlsx"))


def _write_two_sheet_book(dest) -> None:
    import datetime

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Costs"
    wb.active.append(["a", "b"])
    wb.active.append([None, None])
    wb.active.append(["c", None])
    notes = wb.create_sheet("Notes")
    notes.append(["Reviewed", datetime.date(2026, 9, 4)])
    wb.save(dest)


def test_xlsx_extract_cells_without_sheet_returns_the_index_not_cells(tmp_path) -> None:
    dest = tmp_path / "book.xlsx"
    _write_two_sheet_book(dest)

    result = server.xlsx_extract_cells(str(dest))

    assert result == {
        "sheets": [
            {"title": "Costs", "dimensions": "A1:B3", "rows": 3, "non_empty_cells": 3},
            {"title": "Notes", "dimensions": "A1:B1", "rows": 1, "non_empty_cells": 2},
        ]
    }


def test_xlsx_extract_cells_with_sheet_returns_rows_keyed_by_excel_row(tmp_path) -> None:
    dest = tmp_path / "book.xlsx"
    _write_two_sheet_book(dest)

    result = server.xlsx_extract_cells(str(dest), sheet="Costs")

    assert result["title"] == "Costs"
    assert result["rows"] == {"1": ["a", "b"], "3": ["c"]}


@pytest.mark.asyncio
async def test_xlsx_extract_cells_tool_result_is_unwrapped_json_in_both_modes(tmp_path) -> None:
    # Through FastMCP's own result conversion (what the wire carries), not the
    # bare function: a Union/Mapping return annotation would come back wrapped
    # as {"result": ...} and a non-JSON value would fail serialisation here.
    import json

    dest = tmp_path / "book.xlsx"
    _write_two_sheet_book(dest)

    unstructured, structured = await server.mcp.call_tool("xlsx_extract_cells", {"src": str(dest)})
    assert list(structured) == ["sheets"]
    assert json.loads(unstructured[0].text)["sheets"][1]["title"] == "Notes"

    unstructured, structured = await server.mcp.call_tool(
        "xlsx_extract_cells", {"src": str(dest), "sheet": "Notes"}
    )
    assert structured["rows"] == {"1": ["Reviewed", "2026-09-04T00:00:00"]}
    assert json.loads(unstructured[0].text)["rows"] == structured["rows"]


@pytest.mark.asyncio
async def test_xlsx_extract_cells_tool_describes_the_two_step_read() -> None:
    # The description IS the agent's contract for this tool; it must say the
    # sheet-less call is an index and how rows are keyed.
    tools = await server.mcp.list_tools()
    description = next(t for t in tools if t.name == "xlsx_extract_cells").description or ""
    assert "index" in description
    assert "row number" in description
    assert "trailing blank" in description


@pytest.mark.asyncio
async def test_office_shell_tool_is_registered() -> None:
    tools = await server.mcp.list_tools()
    names = {t.name for t in tools}
    assert "office_shell" in names


@pytest.mark.asyncio
async def test_office_shell_tool_describes_cmd_and_cwd() -> None:
    tools = await server.mcp.list_tools()
    tool = next(t for t in tools if t.name == "office_shell")
    schema = tool.inputSchema
    assert "cmd" in schema["properties"]
    assert "cwd" in schema["properties"]
    assert "cmd" in schema.get("required", [])


def test_office_shell_empty_command_raises() -> None:
    from src.office_shell import OfficeShellError

    with pytest.raises(OfficeShellError):
        server.office_shell("   ")


def test_office_shell_runs_a_real_command(tmp_path) -> None:
    result = server.office_shell("echo hello", cwd=str(tmp_path))
    assert result["exit_code"] == 0
    assert result["stdout"].strip() == "hello"
    assert result["timed_out"] is False


def test_xlsx_extract_cells_tool_passes_the_row_window_through(tmp_path: Path) -> None:
    """Observes, if absent: min_row/max_row accepted by the tool signature but
    dropped before the extractor (all rows come back)."""
    from openpyxl import Workbook

    from src import server

    wb = Workbook()
    ws = wb.active
    ws.title = "Rows"
    for i in range(1, 6):
        ws.append([i])
    dest = tmp_path / "rows.xlsx"
    wb.save(dest)

    out = server.xlsx_extract_cells(str(dest), sheet="Rows", min_row=2, max_row=3)
    assert list(out["rows"]) == ["2", "3"]
    # FastMCP's @mcp.tool() registers and returns the function itself; its
    # docstring is the description the agent reads.
    description = server.xlsx_extract_cells.__doc__ or ""
    assert "min_row" in description and "absolute" in description.lower()
