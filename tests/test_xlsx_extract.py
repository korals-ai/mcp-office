"""Tests for the two-step Excel read: ``list_sheets`` (index) and
``extract_sheet`` (one sheet, rows keyed by Excel row number).

Pure-python (openpyxl), so this runs a real round-trip: write a workbook
with openpyxl, read it back, assert on the values — no soffice/poppler needed.
"""

from __future__ import annotations

import datetime
import json
from collections.abc import Callable
from pathlib import Path

import pytest

from src.xlsx_extract import XlsxExtractError, extract_sheet, list_sheets

# Rows 1-4 as the original tool wrote them; row 3 is wholly blank so row
# numbering and empty-row omission are both exercised by one sheet.
COSTSHEET_ROWS: list[list[object]] = [
    ["P/N", "Qty", "Unit Price"],
    ["PNM-C32083RVQ", 4, 121.4],
    [None, None, None],
    ["1LAN-SFP-4305BC-U", 5, 0],
]


def _write_workbook(dest: Path, sheets: dict[str, list[list[object]]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for title, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        for row in rows:
            ws.append(row)
    wb.save(dest)


@pytest.fixture
def book(tmp_path: Path) -> Path:
    src = tmp_path / "book.xlsx"
    _write_workbook(
        src,
        {
            "Costsheet": COSTSHEET_ROWS,
            "Notes": [["Reviewed", datetime.date(2026, 9, 4)]],
        },
    )
    return src


def _read_any(src: Path) -> object:
    return extract_sheet(src, "whatever")


@pytest.mark.parametrize("read", [list_sheets, _read_any])
def test_missing_source_raises(tmp_path: Path, read: Callable[[Path], object]) -> None:
    with pytest.raises(XlsxExtractError, match="source does not exist"):
        read(tmp_path / "nope.xlsx")


def test_directory_source_raises(tmp_path: Path) -> None:
    d = tmp_path / "adir"
    d.mkdir()
    with pytest.raises(XlsxExtractError, match="source is not a file"):
        list_sheets(d)


def test_corrupt_file_raises(tmp_path: Path) -> None:
    src = tmp_path / "bad.xlsx"
    src.write_bytes(b"not a real xlsx")
    with pytest.raises(XlsxExtractError, match="could not open"):
        list_sheets(src)


# --- step 1: the index -----------------------------------------------------


def test_index_lists_sheets_in_workbook_order_with_counts_and_no_cell_data(book: Path) -> None:
    index = list_sheets(book)

    assert index == {
        "sheets": [
            # 4 rows in the used range, 3+3+0+3 = 9 values (a real 0 counts).
            {"title": "Costsheet", "dimensions": "A1:C4", "rows": 4, "non_empty_cells": 9},
            {"title": "Notes", "dimensions": "A1:B1", "rows": 1, "non_empty_cells": 2},
        ]
    }


def test_index_is_not_bounded_by_max_cells(book: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # The index must be answerable for ANY workbook — it is the step the
    # agent takes to decide which sheet fits, so the ceiling cannot apply.
    monkeypatch.setattr("src.xlsx_extract.MAX_CELLS", 1)

    assert [s["title"] for s in list_sheets(book)["sheets"]] == ["Costsheet", "Notes"]


# --- step 2: one sheet -----------------------------------------------------


def test_sheet_rows_keyed_by_excel_row_number_with_typed_values(book: Path) -> None:
    result = extract_sheet(book, "Costsheet")

    assert result["title"] == "Costsheet"
    assert result["dimensions"] == "A1:C4"
    assert result["rows"] == {
        "1": ["P/N", "Qty", "Unit Price"],
        "2": ["PNM-C32083RVQ", 4, 121.4],
        # row 3 is wholly blank: omitted, and row 4 keeps its Excel number
        "4": ["1LAN-SFP-4305BC-U", 5, 0],  # a real 0, not blank
    }


def test_trailing_blanks_trimmed_but_interior_and_leading_blanks_kept(tmp_path: Path) -> None:
    src = tmp_path / "book.xlsx"
    _write_workbook(
        src,
        {
            "S": [
                ["a", None, "c", None, None],
                [None, None, "only-in-C"],
                ["b", None, None],
            ]
        },
    )

    rows = extract_sheet(src, "S")["rows"]

    assert rows["1"] == ["a", None, "c"]  # interior null kept, trailing pair gone
    assert rows["2"] == [None, None, "only-in-C"]  # column C stays at index 2
    assert rows["3"] == ["b"]


def test_empty_rows_omitted_while_row_numbers_stay_excel_accurate(tmp_path: Path) -> None:
    src = tmp_path / "book.xlsx"
    blank: list[object] = [None]
    _write_workbook(
        src,
        {"S": [["r1"], blank, blank, ["r4"], blank, blank, ["r7"]]},
    )

    rows = extract_sheet(src, "S")["rows"]

    assert list(rows) == ["1", "4", "7"]
    assert rows["7"] == ["r7"]


def test_far_column_only_content_returns_that_row_key_positional_from_a(tmp_path: Path) -> None:
    from openpyxl import Workbook

    src = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "S"
    wb.active["H7"] = "x"
    wb.save(src)

    index = list_sheets(src)["sheets"][0]
    assert index == {"title": "S", "dimensions": "H7:H7", "rows": 1, "non_empty_cells": 1}

    result = extract_sheet(src, "S")
    assert result["dimensions"] == "H7:H7"
    # Still keyed by the Excel row and positional from column A (H = index 7),
    # so the agent's rows["7"][7] addressing works whatever the used range is.
    assert result["rows"] == {"7": [None] * 7 + ["x"]}


def test_formula_reads_as_cached_result_never_formula_text(tmp_path: Path) -> None:
    src = tmp_path / "book.xlsx"
    _write_workbook(src, {"S": [[2, 3, "=A1*B1"], ["=A1+B1"]]})

    rows = extract_sheet(src, "S")["rows"]

    # openpyxl writes no cached result, so with data_only=True both formulas
    # read as blank: the trailing one in row 1 is trimmed, row 2 vanishes.
    assert rows == {"1": [2, 3]}
    assert "=A1*B1" not in json.dumps(rows)


def test_both_shapes_are_json_serialisable_including_dates(book: Path) -> None:
    index = list_sheets(book)
    cells = extract_sheet(book, "Notes")

    # json.dumps raises TypeError on a datetime — this is what would fail if
    # date coercion were dropped.
    assert json.loads(json.dumps(index)) == index
    assert json.loads(json.dumps(cells))["rows"] == {"1": ["Reviewed", "2026-09-04T00:00:00"]}


def test_unknown_sheet_name_raises_with_available_list(book: Path) -> None:
    with pytest.raises(XlsxExtractError, match=r"not found.*Costsheet.*Notes"):
        extract_sheet(book, "Nope")


def test_oversized_sheet_raises_naming_one_sheet_and_smaller_sheets_still_read(
    book: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # The ceiling is per sheet (A1..last cell rectangle): Costsheet is 4x3=12
    # cells, Notes is 1x2=2. At a ceiling of 4 one is over and one is not,
    # which a workbook-wide count (14) would get wrong for Notes.
    monkeypatch.setattr("src.xlsx_extract.MAX_CELLS", 4)

    with pytest.raises(XlsxExtractError, match=r"'Costsheet'.*12 cells.*ONE sheet"):
        extract_sheet(book, "Costsheet")

    assert extract_sheet(book, "Notes")["rows"]["1"] == ["Reviewed", "2026-09-04T00:00:00"]


def test_chart_sheets_are_skipped_by_the_index_and_named_as_missing(tmp_path: Path) -> None:
    """A workbook with a chart sheet used to raise a raw AttributeError on the
    very first (index) call. Observes, if absent: list_sheets raising, or the
    chart sheet appearing in the index / being readable as cells."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["a", 1])
    ws.append(["b", 2])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2))
    # An EMPTY chart sheet does not survive openpyxl's own save/load; a real
    # one (chart attached) is what a user's workbook carries anyway.
    wb.create_chartsheet("Chart1").add_chart(chart)
    dest = tmp_path / "charts.xlsx"
    wb.save(dest)

    index = list_sheets(dest)
    assert [s["title"] for s in index["sheets"]] == ["Data"]
    with pytest.raises(XlsxExtractError, match=r"not found.*available: \['Data'\]"):
        extract_sheet(dest, "Chart1")


def test_index_counts_from_the_sparse_store_not_the_rectangle(tmp_path: Path) -> None:
    """One value far down a sheet must not cost a walk of the whole A1 rectangle.
    Observes, if absent: a wrong non_empty_cells count (the rectangle walk
    counts the same cells, so this pins the count; the memory bound is the
    reason for _cells and is documented there)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sparse"
    ws["A1"] = "hdr"
    ws["Z5000"] = 1
    dest = tmp_path / "sparse.xlsx"
    wb.save(dest)

    index = list_sheets(dest)
    assert index["sheets"][0]["non_empty_cells"] == 2
    assert index["sheets"][0]["dimensions"] == "A1:Z5000"


def test_row_window_pages_with_excel_accurate_keys(tmp_path: Path) -> None:
    """Observes, if absent: min_row/max_row ignored (all rows back), or keys
    renumbered from 1 inside the window instead of the sheet's real numbers."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Big"
    for i in range(1, 11):
        ws.append([f"r{i}", i])
    dest = tmp_path / "big.xlsx"
    wb.save(dest)

    page = extract_sheet(dest, "Big", min_row=4, max_row=6)
    assert list(page["rows"]) == ["4", "5", "6"]
    assert page["rows"]["5"] == ["r5", 5]
    tail = extract_sheet(dest, "Big", min_row=9)
    assert list(tail["rows"]) == ["9", "10"]
    with pytest.raises(XlsxExtractError, match="bad row window"):
        extract_sheet(dest, "Big", min_row=6, max_row=4)


def test_row_window_bounds_the_ceiling_not_the_whole_sheet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A sheet over the ceiling is still readable a window at a time.
    Observes, if absent: the windowed call refusing because the WHOLE sheet
    is over the ceiling."""
    from openpyxl import Workbook

    from src import xlsx_extract

    wb = Workbook()
    ws = wb.active
    ws.title = "Wide"
    for i in range(1, 7):
        ws.append([i, i, i])
    dest = tmp_path / "wide.xlsx"
    wb.save(dest)
    monkeypatch.setattr(xlsx_extract, "MAX_CELLS", 9)

    with pytest.raises(XlsxExtractError, match="page it with min_row/max_row"):
        extract_sheet(dest, "Wide")
    page = extract_sheet(dest, "Wide", min_row=1, max_row=3)
    assert list(page["rows"]) == ["1", "2", "3"]


def test_legacy_xls_is_refused_with_the_convert_tool_named(tmp_path: Path) -> None:
    """openpyxl's own message says "use xlrd", a library the agent's pod lacks.
    Observes, if absent: the xlrd text reaching the agent, or an .xls being
    opened as a zip and failing with a corrupt-file message."""
    dest = tmp_path / "old.xls"
    dest.write_bytes(b"\xd0\xcf\x11\xe0" + b"\0" * 64)

    with pytest.raises(XlsxExtractError) as excinfo:
        list_sheets(dest)
    assert "mcp__workspace-tool-office__convert" in str(excinfo.value)
    assert "xlrd" not in str(excinfo.value)


def test_duration_cells_read_as_seconds(tmp_path: Path) -> None:
    """A [h]:mm:ss cell comes back from openpyxl as a timedelta.
    Observes, if absent: json.dumps raising, or an ISO duration string the
    agent cannot sum."""
    import json

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Hours"
    ws["A1"] = 1.5
    ws["A1"].number_format = "[h]:mm:ss"
    dest = tmp_path / "hours.xlsx"
    wb.save(dest)

    cells = extract_sheet(dest, "Hours")
    assert cells["rows"]["1"] == [1.5 * 86400]
    json.dumps(cells)
